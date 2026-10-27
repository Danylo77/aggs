from docx import Document
from docx.oxml.ns import qn
d=Document('Практика_Звіт_Теплий.docx')
for i,p in enumerate(d.paragraphs):
 br=p._p.xpath('.//w:br')
 sect=p._p.xpath('.//w:sectPr')
 if br or sect or p.style.name.startswith('Heading 1'):
  print(i,repr(p.text[:80]),p.style.name,'breaks',[x.get(qn('w:type')) for x in br],'sect',bool(sect))
for i in range(15,42):
 p=d.paragraphs[i]
 print(i,repr(p.text),p.style.name)
