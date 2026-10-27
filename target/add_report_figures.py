from pathlib import Path
import math
import numpy as np
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BASE = Path('Практика_Звіт_Теплий_повний.docx')
OUTPUT = Path('Практика_Звіт_Теплий_повний_з_рисунками.docx')
FIGURES = Path('report_figures')
FIGURES.mkdir(exist_ok=True)
plt.rcParams.update({'font.family': 'DejaVu Sans', 'font.size': 9})


def figure_51():
    fig, ax = plt.subplots(figsize=(12, 6.3))
    ax.set_xlim(0, 12); ax.set_ylim(0, 7); ax.axis('off')
    def box(x,y,w,h,text,color):
        p=FancyBboxPatch((x,y),w,h,boxstyle='round,pad=0.04,rounding_size=0.08',
                         facecolor=color,edgecolor='#31465a',linewidth=1.3)
        ax.add_patch(p); ax.text(x+w/2,y+h/2,text,ha='center',va='center',fontsize=9,weight='bold',wrap=True)
        return (x,y,w,h)
    def arrow(a,b,label=None):
        x1=a[0]+a[2]; y1=a[1]+a[3]/2; x2=b[0]; y2=b[1]+b[3]/2
        ax.annotate('',xy=(x2,y2),xytext=(x1,y1),arrowprops=dict(arrowstyle='->',lw=1.4,color='#334e68'))
        if label: ax.text((x1+x2)/2,(y1+y2)/2+0.18,label,ha='center',fontsize=8,color='#334e68')
    b1=box(.3,5.1,1.5,.85,'WKT + цільове\nскорочення','#d9eaf7')
    b2=box(2.15,5.1,1.55,.85,'AdaptiveSimplifier\nвалідація входу','#d9eaf7')
    b3=box(4.05,5.1,1.55,.85,'GeometryParts\nдекомпозиція','#e7f2dd')
    b4=box(5.95,5.1,2.0,.85,'ComplexityAnalyzer\nлокальні ознаки','#fff0cc')
    b5=box(8.3,5.1,1.55,.85,'ToleranceModel\nполе допусків','#ffe3c7')
    b6=box(10.2,5.1,1.5,.85,'VariableToleranceDP\nранжування','#ffd9d9')
    for a,b in zip([b1,b2,b3,b4,b5],[b2,b3,b4,b5,b6]): arrow(a,b)
    bc=box(6.15,3.55,1.6,.72,'BoxCounting\nD та R²','#fff6df')
    gf=box(4.15,3.55,1.6,.72,'Кут, відхилення,\nщільність','#fff6df')
    ax.annotate('',xy=(6.85,5.1),xytext=(6.85,4.27),arrowprops=dict(arrowstyle='->',lw=1.2,color='#8a6d1d'))
    ax.annotate('',xy=(6.15,5.1),xytext=(4.95,4.27),arrowprops=dict(arrowstyle='->',lw=1.2,color='#8a6d1d'))
    eng=box(7.25,2.1,2.2,.85,'AdaptiveEngine\nпоріг, retry, fallback','#eadcf8')
    valb=box(9.85,2.1,1.8,.85,'GeometryValidator\nструктура + JTS validity','#eadcf8')
    out=box(9.85,.55,1.8,.75,'Результат WKT\n+ діагностика','#d9eaf7')
    ax.annotate('',xy=(8.35,2.95),xytext=(10.95,5.1),arrowprops=dict(arrowstyle='->',lw=1.4,color='#334e68'))
    arrow(eng,valb,'кандидат')
    ax.annotate('менше вилучень',xy=(8.3,2.15),xytext=(10.7,1.65),ha='center',fontsize=8,color='#8a3d3d',
                arrowprops=dict(arrowstyle='->',connectionstyle='arc3,rad=-.35',color='#8a3d3d'))
    ax.annotate('',xy=(10.75,1.3),xytext=(10.75,2.1),arrowprops=dict(arrowstyle='->',lw=1.4,color='#334e68'))
    bench=box(.7,.65,4.6,1.35,'Окремий benchmark-модуль\nAdaptive public API + JTS Douglas–Peucker\nметрики якості, CSV, консольний звіт','#eeeeee')
    ax.annotate('',xy=(3.0,2.0),xytext=(2.9,5.1),arrowprops=dict(arrowstyle='->',lw=1.1,linestyle='--',color='#666'))
    ax.text(6,6.65,'Архітектура адаптивної системи спрощення',ha='center',fontsize=14,weight='bold')
    ax.text(6,.1,'Production-алгоритм і експериментальна інфраструктура розділені',ha='center',fontsize=8,color='#555')
    fig.tight_layout(); fig.savefig(FIGURES/'figure_5_1_architecture.png',dpi=260,bbox_inches='tight'); plt.close(fig)


def figure_52():
    n=105
    steps=np.ones(n); steps[12:30]=.35; steps[68:82]=.45
    x=np.cumsum(steps)-steps[0]
    y=.12*np.sin(x/7)
    y[40:65] += np.array([(-1)**i*(.45+.025*i) for i in range(25)])
    y[65:] += .25*np.sin((x[65:]-x[65])/2.2)
    angle=np.zeros(n); dev=np.zeros(n); dens=np.zeros(n)
    seg=np.hypot(np.diff(x),np.diff(y)); meanseg=seg.mean()
    for i in range(1,n-1):
        u=np.array([x[i]-x[i-1],y[i]-y[i-1]]); v=np.array([x[i+1]-x[i],y[i+1]-y[i]])
        angle[i]=math.acos(np.clip(np.dot(u,v)/(np.linalg.norm(u)*np.linalg.norm(v)),-1,1))/math.pi
    h=2
    for i in range(n):
        lo=max(0,i-h); hi=min(n-1,i+h); pts=np.column_stack([x[lo:hi+1],y[lo:hi+1]])
        path=np.hypot(np.diff(pts[:,0]),np.diff(pts[:,1])).sum()
        a=pts[0]; b=pts[-1]; ab=b-a; den=np.dot(ab,ab)
        distances=[]
        for p in pts[1:-1]:
            t=np.clip(np.dot(p-a,ab)/den,0,1) if den else 0
            distances.append(np.linalg.norm(p-(a+t*ab)))
        dev[i]=min(1,2*max(distances,default=0)/path) if path else 0
        local=path/max(1,hi-lo); r=meanseg/local if local else 1
        dens[i]=r/(1+r)
    complexity=(.35*angle+.30*dev+.10*dens)/.75
    raw=np.log(4)*(1-2*complexity); smooth=np.empty(n)
    for i in range(n): smooth[i]=np.mean(raw[max(0,i-2):min(n,i+3)])
    mult=np.exp(smooth)
    fig,axes=plt.subplots(3,1,figsize=(12,7),sharex=True,gridspec_kw={'height_ratios':[1.2,1,1]})
    axes[0].plot(x,y,color='#3b6ea8',lw=1.5); axes[0].scatter(x,y,s=9,color='#3b6ea8')
    axes[0].set_ylabel('Геометрія'); axes[0].set_title('Ілюстративний вплив локальних ознак на допуск',weight='bold',fontsize=13)
    axes[1].plot(x,complexity,color='#c44e52',lw=1.8); axes[1].fill_between(x,0,complexity,color='#c44e52',alpha=.18)
    axes[1].set_ylabel('Складність cᵢ'); axes[1].set_ylim(0,1)
    axes[2].plot(x,mult,color='#2b8a3e',lw=1.8); axes[2].axhline(1,color='#777',lw=.8,ls='--')
    axes[2].set_ylabel('Множник mᵢ'); axes[2].set_xlabel('Положення вздовж лінії (вхідні координатні одиниці)')
    for ax in axes:
        ax.axvspan(x[40],x[64],color='#f5b7b1',alpha=.2); ax.grid(alpha=.22)
    axes[1].annotate('локально складний фрагмент',xy=(x[53],complexity[53]),xytext=(x[66],.83),
                     arrowprops=dict(arrowstyle='->',color='#8a1c1c'),color='#8a1c1c')
    axes[2].annotate('менший множник → обережніше спрощення',xy=(x[53],mult[53]),xytext=(x[61],3.1),
                     arrowprops=dict(arrowstyle='->',color='#1f6f32'),color='#1f6f32')
    axes[2].annotate('простіша ділянка → більший допуск',xy=(x[18],mult[18]),xytext=(x[4],1.1),
                     arrowprops=dict(arrowstyle='->',color='#1f6f32'),color='#1f6f32')
    fig.text(.5,.015,'Профіль обчислено за геометричними ознаками fallback; mᵢ = exp(mean(ln 4 · (1 − 2cᵢ)))',ha='center',fontsize=8,color='#555')
    fig.tight_layout(rect=[0,.035,1,1]); fig.savefig(FIGURES/'figure_5_2_tolerance_profile.png',dpi=260,bbox_inches='tight'); plt.close(fig)


def figure_61():
    labels=['Hausdorff','Середнє\nвідхилення','Зміна\nдовжини','Зміна\nплощі*','Симетрична\nрізниця*','Макс. відхилення\nвершин']
    adaptive=np.array([18,42,77,32,15,18]); dp=np.array([126,124,77,68,88,126]); ties=np.array([69,47,59,12,9,69])
    x=np.arange(len(labels)); width=.25
    fig,ax=plt.subplots(figsize=(12,6.3))
    bars=[ax.bar(x-width,adaptive,width,label='Adaptive',color='#3b82b6'),
          ax.bar(x,dp,width,label='Douglas–Peucker',color='#e07a2d'),
          ax.bar(x+width,ties,width,label='Однаково',color='#9aa0a6')]
    for group in bars:
        ax.bar_label(group,padding=2,fontsize=8)
    ax.set_xticks(x,labels); ax.set_ylabel('Кількість узгоджених порівнянь')
    ax.set_title('Результати за метриками якості',fontsize=14,weight='bold')
    ax.legend(ncol=3,loc='upper center'); ax.grid(axis='y',alpha=.25); ax.set_ylim(0,145)
    ax.text(.5,-.16,'* Метрики площі застосовні лише до 112 узгоджених полігональних порівнянь; інші — до 213.',
            transform=ax.transAxes,ha='center',fontsize=8,color='#555')
    fig.tight_layout(); fig.savefig(FIGURES/'figure_6_1_metric_wins.png',dpi=260,bbox_inches='tight'); plt.close(fig)


def set_font(run, size=14, bold=False):
    run.font.name='Times New Roman'; run._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman')
    run.font.size=Pt(size); run.bold=bold; run.font.color.rgb=RGBColor(0,0,0)


def insert_caption_after(paragraph, text):
    new_p=OxmlElement('w:p'); paragraph._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    p=Paragraph(new_p, paragraph._parent)
    p.style='Normal'; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=0; p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(8)
    r=p.add_run(text); set_font(r,14)
    return p


def replace_marker(doc, needle, image, caption, alt):
    p=next(p for p in doc.paragraphs if needle in p.text and 'ПОТРІБНО ДОДАТИ РИСУНОК' in p.text)
    p.clear(); p.style='Normal'; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=0; p.paragraph_format.keep_with_next=True
    run=p.add_run(); inline=run.add_picture(str(image),width=Inches(6.45))
    doc_pr=inline._inline.docPr; doc_pr.set('descr',alt); doc_pr.set('title',caption)
    insert_caption_after(p,caption)


figure_51(); figure_52(); figure_61()
doc=Document(BASE)
replace_marker(doc,'Рисунок 5.1',FIGURES/'figure_5_1_architecture.png',
               'Рисунок 5.1 – Архітектура адаптивної системи спрощення геометрії',
               'Схема компонентів адаптивного спрощення та окремого benchmark-модуля')
replace_marker(doc,'Рисунок 5.2',FIGURES/'figure_5_2_tolerance_profile.png',
               'Рисунок 5.2 – Вплив локальної складності на множник адаптивного допуску',
               'Ілюстративний профіль геометрії, локальної складності та множника допуску')
replace_marker(doc,'Рисунок 6.1',FIGURES/'figure_6_1_metric_wins.png',
               'Рисунок 6.1 – Кількість кращих результатів за основними метриками якості',
               'Групована діаграма перемог Adaptive, Douglas–Peucker та однакових результатів')
# Keep field refresh request.
settings=doc.settings.element; update=settings.find(qn('w:updateFields'))
if update is None: update=OxmlElement('w:updateFields'); settings.append(update)
update.set(qn('w:val'),'true')
doc.save(OUTPUT)

check=Document(OUTPUT)
assert len(check.inline_shapes)==4
assert sum('ПОТРІБНО ДОДАТИ РИСУНОК' in p.text for p in check.paragraphs)==5
for caption in ('Рисунок 5.1','Рисунок 5.2','Рисунок 6.1'):
    assert any(p.text.startswith(caption) and p.alignment==WD_ALIGN_PARAGRAPH.CENTER for p in check.paragraphs)
print(f'WROTE {OUTPUT}; figures={len(check.inline_shapes)-1}; remaining_placeholders=5')
