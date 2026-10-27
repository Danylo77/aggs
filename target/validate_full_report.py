from docx import Document
from docx.oxml.ns import qn
from pathlib import Path
src=Document('Практика_Звіт_Теплий.docx')
out=Document('Практика_Звіт_Теплий_повний.docx')
print('size_bytes',Path('Практика_Звіт_Теплий_повний.docx').stat().st_size)
print('paragraphs',len(out.paragraphs),'tables',len(out.tables),'sections',len(out.sections),'images',len(out.inline_shapes))
print('headings')
for p in out.paragraphs:
 if p.style.name.startswith('Heading') and (p.text.startswith('Розділ 5') or p.text.startswith('Розділ 6') or p.text.startswith('5.') or p.text.startswith('6.') or p.text.strip()=='Висновок'):
  print(p.style.name, p.text.strip())
markers=[p for p in out.paragraphs if 'ПОТРІБНО ДОДАТИ РИСУНОК' in p.text]
print('markers',len(markers),'all_red',all(any(r.font.color.rgb and str(r.font.color.rgb)=='FF0000' for r in p.runs) for p in markers))
print('update_fields',bool(out.settings.element.xpath('.//w:updateFields[@w:val="true"]')))
print('toc_fields',len(out.element.body.xpath('.//w:instrText[contains(text(),"TOC")]')))
print('page_fields_body',len(out.element.body.xpath('.//w:instrText[contains(text(),"PAGE")]')))
print('footer_fields',sum(len(s.footer._element.xpath('.//w:instrText[contains(text(),"PAGE")]')) for s in out.sections))
print('margins_equal',all((a.page_width,a.page_height,a.top_margin,a.right_margin,a.bottom_margin,a.left_margin,a.header_distance,a.footer_distance)==(b.page_width,b.page_height,b.top_margin,b.right_margin,b.bottom_margin,b.left_margin,b.header_distance,b.footer_distance) for a,b in zip(src.sections,out.sections)))
print('theory_prefix_equal',all(a.text==b.text and a.style.name==b.style.name for a,b in zip(src.paragraphs[:137],out.paragraphs[:137])))
for i,t in enumerate(out.tables,1): print('table',i,'rows',len(t.rows),'cols',len(t.columns))
