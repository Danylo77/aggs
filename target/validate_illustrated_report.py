from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path
from zipfile import ZipFile
base=Document('Практика_Звіт_Теплий_повний.docx')
out=Document('Практика_Звіт_Теплий_повний_з_рисунками.docx')
print('size_bytes',Path('Практика_Звіт_Теплий_повний_з_рисунками.docx').stat().st_size)
print('paragraphs',len(out.paragraphs),'tables',len(out.tables),'inline_shapes',len(out.inline_shapes),'remaining_markers',sum('ПОТРІБНО ДОДАТИ РИСУНОК' in p.text for p in out.paragraphs))
print('margins_equal',all((a.page_width,a.page_height,a.top_margin,a.right_margin,a.bottom_margin,a.left_margin)==(b.page_width,b.page_height,b.top_margin,b.right_margin,b.bottom_margin,b.left_margin) for a,b in zip(base.sections,out.sections)))
print('footer_equal', [p.text for s in base.sections for p in s.footer.paragraphs]==[p.text for s in out.sections for p in s.footer.paragraphs])
print('theory_equal',all(a.text==b.text and a.style.name==b.style.name for a,b in zip(base.paragraphs[:137],out.paragraphs[:137])))
print('update_fields',bool(out.settings.element.xpath('.//w:updateFields[@w:val="true"]')))
for key in ('Рисунок 5.1','Рисунок 5.2','Рисунок 6.1'):
 ps=[p for p in out.paragraphs if p.text.startswith(key)]
 print(key,'captions',len(ps),'centered',all(p.alignment==WD_ALIGN_PARAGRAPH.CENTER for p in ps))
with ZipFile('Практика_Звіт_Теплий_повний_з_рисунками.docx') as z:
 print('embedded_media',[n for n in z.namelist() if n.startswith('word/media/')])
