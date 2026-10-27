from docx import Document
from pathlib import Path
p=Path('Практика_Звіт_Теплий.docx')
d=Document(p)
print(f'paragraphs={len(d.paragraphs)} tables={len(d.tables)} sections={len(d.sections)} inline_shapes={len(d.inline_shapes)}')
for si,s in enumerate(d.sections):
 print(f'SECTION {si}: page={s.page_width}x{s.page_height} margins={s.top_margin},{s.right_margin},{s.bottom_margin},{s.left_margin} header={s.header_distance} footer={s.footer_distance}')
 print(' HEADER:', ' | '.join(x.text for x in s.header.paragraphs if x.text.strip()))
 print(' FOOTER:', ' | '.join(x.text for x in s.footer.paragraphs if x.text.strip()))
print('STYLES USED:')
from collections import Counter
print(Counter(p.style.name for p in d.paragraphs))
print('PARAGRAPHS:')
for i,p in enumerate(d.paragraphs):
 t=' '.join(p.text.split())
 if t:
  fmt=p.paragraph_format
  run=p.runs[0] if p.runs else None
  print(f'{i:04d}\t{p.style.name}\talign={p.alignment}\tbefore={fmt.space_before}\tafter={fmt.space_after}\tline={fmt.line_spacing}\tindent={fmt.first_line_indent}\tfont={(run.font.name if run else None)}/{(run.font.size.pt if run and run.font.size else None)}/{(run.bold if run else None)}\t{t[:500]}')
print('TABLES:')
for ti,t in enumerate(d.tables):
 print(f'TABLE {ti}: rows={len(t.rows)} cols={len(t.columns)} style={t.style.name if t.style else None}')
 for r in t.rows[:4]: print(' | '.join(' '.join(c.text.split())[:100] for c in r.cells))
